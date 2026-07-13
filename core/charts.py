"""
charts.py
------------------------------------
Create Excel Charts
"""

from openpyxl.chart import (
    BarChart,
    PieChart,
    LineChart,
    Reference
)


def create_region_chart(sheet, start_row, end_row):
    """
    Create region sales bar chart.
    """

    data = Reference(
        sheet,
        min_col=2,
        min_row=start_row,
        max_row=end_row
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=start_row,
        max_row=end_row
    )

    chart = BarChart()

    chart.add_data(
        data,
        titles_from_data=False
    )

    chart.set_categories(categories)

    chart.title = "Region-wise Sales"

    chart.y_axis.title = "Sales"

    chart.x_axis.title = "Region"

    sheet.add_chart(
        chart,
        "D3"
    )


def create_product_chart(sheet, start_row, end_row):
    """
    Create product sales pie chart.
    """

    data = Reference(
        sheet,
        min_col=2,
        min_row=start_row,
        max_row=end_row
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=start_row,
        max_row=end_row
    )

    chart = PieChart()

    chart.add_data(
        data,
        titles_from_data=False
    )

    chart.set_categories(categories)

    chart.title = "Product-wise Sales"

    sheet.add_chart(
        chart,
        "D20"
    )


def create_monthly_chart(sheet, start_row, end_row):
    """
    Create monthly sales trend chart.
    """

    data = Reference(
        sheet,
        min_col=2,
        min_row=start_row,
        max_row=end_row
    )

    categories = Reference(
        sheet,
        min_col=1,
        min_row=start_row,
        max_row=end_row
    )

    chart = LineChart()

    chart.add_data(data)

    chart.set_categories(categories)

    chart.title = "Monthly Sales Trend"

    chart.y_axis.title = "Sales"

    chart.x_axis.title = "Month"

    sheet.add_chart(
        chart,
        "D37"
    )